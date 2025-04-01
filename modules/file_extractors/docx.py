# modules/file_extractors/docx.py
"""
DOCX file content extractor.
"""
import os
import asyncio
from typing import Optional

from modules.file_extractors.base import BaseFileExtractor
from modules.utils.logging import get_logger

logger = get_logger(__name__)

class DocxFileExtractor(BaseFileExtractor):
    """
    Extractor for DOCX files.
    """
    
    async def extract(self, file_path: str) -> Optional[str]:
        """
        Extract content from a DOCX file.
        
        Args:
            file_path: Path to the file
            
        Returns:
            File content as string or None if extraction failed
        """
        try:
            # Use asyncio to read the file without blocking
            content = await asyncio.to_thread(self._extract_docx, file_path)
            return content
        except Exception as e:
            logger.error(f"Error extracting text from DOCX {file_path}: {str(e)}")
            return None
    
    def _extract_docx(self, file_path: str) -> str:
        """
        Extract text from a DOCX file.
        
        Args:
            file_path: Path to the DOCX file
            
        Returns:
            Extracted text
        """
        import docx
        
        doc = docx.Document(file_path)
        full_text = []
        
        # Extract text from paragraphs
        for para in doc.paragraphs:
            full_text.append(para.text)
            
        # Extract text from tables
        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    full_text.append(cell.text)
        
        return '\n\n'.join(full_text)